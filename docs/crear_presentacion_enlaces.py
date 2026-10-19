"""Genera la presentación y su versión textual con python-pptx."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = Path(__file__).resolve().parent
slides = [
('Enlaces redundantes en i4a', 'Qué módulos cambiaron y por qué', [
'Comparación: CoNexDat/i4a_project 8df7996 → feat/redundant-link-discovery 81aa146.',
'Objetivo: descubrir un vecino alternativo y recuperar conectividad sin cambiar las IP del nodo ni de sus clientes.',
'Alcance: implementación integrada hasta el 12/10/2026. La instrumentación soak del 15/10 se distingue del cambio funcional.'], 'docs/comnetar-integration.md; docs/three-node-soak.md'),
('El problema que resuelve', 'Una segunda conexión necesita coordinación y rutas de ida y retorno', [
'Un nodo contiene cinco ESP: central y cuatro radios periféricas, coordinados mediante el anillo SPI.',
'Topología buscada: Root–A, Root–B y un lateral A–B. Si cae Root–A, el recorrido esperado es A → B → Root.',
'Conservar una asociación Wi-Fi no alcanza: hay que admitir la sesión, actualizar las tablas y reconstruir el retorno hacia A.',
'La tabla mantiene un único gateway activo; la alternativa se selecciona con el protocolo de búsqueda existente.'], 'docs/redundant-link-discovery.md; docs/gateway-route-recovery.md'),
('Mapa de módulos', 'Cambios funcionales y soporte necesario', [
'Nuevos: discovery_manager, neighbor_manager y peer_session; offer_policy y point_to_point separan reglas verificables.',
'Modificados: node, station/device, client/server, callbacks, channel_manager y routing_hooks.',
'Ruteo y coordinación: routing, sync, shared_state y ring_share.',
'Soporte: node_ota/reset_manager, SPI/netif, configuración RSSI, telemetría y pruebas.'], 'git diff --name-status 8df7996 81aa146 -- components'),
('Descubrimiento: una radio por vez', 'Nuevo discovery_manager + cambios en node y station', [
'Problema: varias radios podían buscar a la vez y competir por la provisión o por el cambio AP/Station.',
'El ESP central arbitra un único turno por SPI. Propietario, sesión y secuencia impiden que mensajes viejos liberen un turno nuevo.',
'node_control concentra los cambios de modo; Station comprueba el permiso antes de escanear y antes de asociarse.',
'Espera AP de 600–700 s; búsqueda hasta 60 s y reserva central de 90 s. Un intento fallido vuelve a AP y libera el turno.'], 'components/integration/discovery_manager/; components/integration/node/node.c; components/wireless/station/station.c'),
('Admisión: conservar el direccionamiento', 'Nueva offer_policy + validación previa al ruteo', [
'Problema: aceptar cualquier oferta podía reprovisionar un nodo ya configurado o admitir una relación incompatible.',
'Se compara provided_network con el bloque lógico original: la oferta redundante debe tener base mayor y prefijo mayor.',
'La regla usa la subred lógica ofrecida, no el /30 de DHCP ni la subred particular de la placa.',
'El primer enlace tiene validación propia. Restaurar el uplink admite su asignación original; rechazar una oferta no renumera el nodo.'], 'components/integration/node/offer_policy.c; components/wireless/client/client.c; docs/redundant-link-discovery.md'),
('Evitar enlaces duplicados', 'Nuevo neighbor_manager + peer_session; cambios en client/server', [
'Problema: A podía conectarse a B por más de una orientación, incluso mediante intentos recíprocos simultáneos.',
'Un registro central por UUID identifica al nodo completo y permite un solo enlace por par de nodos.',
'peer_session intercambia identidad y ejecuta reserva, compromiso y activación antes de entregar eventos al ruteo.',
'Tokens, revisiones y vencimientos protegen sesiones nuevas frente a liberaciones antiguas. El framing TCP reconstruye mensajes fragmentados.'], 'components/integration/neighbor_manager/; components/wireless/peer_session/; components/wireless/client/; components/wireless/server/'),
('Recuperar salida y retorno', 'Cambios en routing: peer, sibling, root y tabla', [
'Al perder una interfaz se retiran sus rutas y se invalida el gateway si dependía de ella; se conservan las subredes existentes.',
'La elección compara DTR completo, evita entradas vacías como rutas por defecto y descarta pedidos que vuelven sobre sí mismos.',
'Las solicitudes se reintentan cada 30 s; Root agenda respuesta a los 10 s sin postergarla por pedidos repetidos.',
'También se inicia recuperación al perder un enlace no predeterminado en radios provisionadas con DTR > 1: permite reconstruir el retorno que Root había retirado.'], 'components/routing/routing/src/forwarder/peer.c; sibling.c; root.c; routing.c; docs/correccion-transporte-vacio-20261012.md'),
('Que el evento de pérdida se procese', 'Cambios en sync, shared_state y colas de routing', [
'Problema: perder una solicitud o un token podía dejar PEER_LOST en cola y conservar una salida ya inválida.',
'SYNC v2 incorpora sesión y secuencia. El central reintenta la misma circulación; un periférico no repite un callback ya completado.',
'Una cola pendiente solicita turno nuevamente. Los eventos concurrentes se separan en lotes protegidos por mutex.',
'SHARED_STATE v2 identifica tablas por sesión, secuencia y propietario. Si publicar falla, retiene el trabajo pendiente y vuelve a publicar.'], 'components/routing/sync/; components/routing/shared_state/; components/routing/routing/src/routing.c; docs/correcciones-coordinacion-v2.md'),
('Evitar desvíos durante el escaneo', 'Cambios en routing_hooks y device/point_to_point', [
'Problema reproducido: una Station sin IP tenía máscara y subred cero; la comparación clasificaba cualquier destino como punto a punto.',
'El hook elegía una Wi-Fi aún no disponible y descartaba tráfico que debía continuar por SPI.',
'Ahora se exige máscara válida y Station con IP, sin cierre en curso. Los destinos ajenos al transporte siguen las reglas normales.',
'Sin gateway se usa una interfaz de descarte con ERR_RTE: devolver NULL permitiría el fallback de lwIP hacia SPI.'], 'components/wireless/device/point_to_point.c; components/integration/routing_hooks/routing_hooks.c; docs/correccion-transporte-vacio-20261012.md'),
('Cambiar de modo sin dejar recursos vivos', 'Cambios en Wi-Fi, TCP y eventos de integración', [
'Station se cancela de forma cooperativa: detiene el escaneo y espera al worker para no eliminar tareas con mutex tomados.',
'Se cierran tareas TCP y se retiran referencias de ruteo antes de destruir interfaces; se ignoran eventos tardíos durante el cierre.',
'El servidor mantiene supervisión y reintenta fallos de socket/bind/listen/select cada 2 s, hasta una solicitud explícita de cierre.',
'callbacks y channel_manager integran la sesión admitida con eventos y canales; una asociación provisional no debe incorporar rutas.'], 'components/wireless/{device,station,access_point,client,server}/; components/integration/{callbacks,channel_manager}/'),
('Coordinar con arranque y OTA', 'Cambios en app_main, node_ota, ring_share y reset_manager', [
'Problema: nuevas búsquedas y tráfico durante mantenimiento pueden competir con OTA o arrancar con versiones incompatibles.',
'La barrera de arranque espera las cinco aplicaciones listas y compatibles; app_main confirma el modo inicial cuando driver y worker están listos.',
'Durante mantenimiento se bloquean nuevas admisiones y turnos. Los callbacks SPI encolan trabajo sin esperar transmisiones de respuesta.',
'STARTUP, SYNC y SHARED_STATE v2 exigen actualización coordinada. No se cambia la tabla de particiones ni se requiere borrar NVS.'], 'components/integration/node_ota/; components/routing/ring_share/; components/integration/reset_manager/; app/main/main.c'),
('Soporte de radio y operación', 'Configuración RSSI y observabilidad', [
'portal_settings, wifi_credentials y config_portal permiten persistir configuración; Station aplica el umbral RSSI por interfaz.',
'antenna_orientation_mode aporta modos especiales. El umbral ayuda a seleccionar candidatos, pero no garantiza eliminar lóbulos traseros.',
'info_manager y los diagnósticos distinguen sesión admitida de asociación Wi-Fi e informan edad de las muestras.',
'Estos cambios apoyan la configuración y la investigación; no sustituyen la admisión por UUID ni la elección de gateway.'], 'components/integration/{portal_settings,wifi_credentials,config_portal,antenna_orientation_mode,info_manager}/; docs/portal-rssi-integration.md'),
('Estabilidad del transporte bajo carga', 'Cambios complementarios en SPI y ring_link_netif', [
'En recepción IP, el callback entrega el pbuf a lwIP o lo libera si se rechaza: evita retener memoria bajo saturación.',
'SPI agrega diagnóstico CRC al recibir DMA y antes de procesar para distinguir corrupción inicial de cambios posteriores.',
'CONFIG_SPI_RX_EXTRA_BUFFERS permite comparar pool base y reservas adicionales sin aumentar los ocho descriptores DMA.',
'Es una variante experimental: extra 8 consume 12.800 bytes adicionales de buffers. No demuestra corrección de CRC ni estabilidad prolongada.'], 'components/internal/spi/; components/internal/ring_link_netif/; docs/manejo-buffers-spi-20261012.md; docs/gateway-route-recovery.md'),
('Qué evidencia existe', 'Separar regresiones de software y resultado físico', [
'Las notas registran pruebas nativas de arbitraje, admisión, ofertas, recuperación, coordinación y servidor, además de pruebas Python y builds ESP-IDF.',
'Coordinación v2 documenta ejecución posterior en RPi de las pruebas antes bloqueadas por Windows.',
'Los cortes físicos del 12/10 mostraron fallos de retorno: seleccionar Sur en la tabla local no acreditó recuperación extremo a extremo.',
'La regresión de pérdida de enlace no predeterminado pasó tras corregir peer.c; esa nota deja pendiente instalar y validar varios cortes.'], 'docs/correcciones-coordinacion-v2.md; docs/correccion-transporte-vacio-20261012.md; tests/run_native_recovery_objects.py'),
('Resultado y siguiente validación', 'La redundancia requiere descubrimiento, admisión y convergencia', [
'Descubrir: una radio libre busca por turno, sin interrumpir enlaces admitidos.',
'Admitir: UUID y oferta lógica evitan duplicados y conservan el direccionamiento.',
'Converger: routing + sync + shared_state retiran rutas inválidas y propagan salida y retorno.',
'Validación pendiente: cortes repetidos en ambos sentidos, IP constantes, tráfico local y WAN; ensayo prolongado con capturas de todas las placas posibles.'], 'docs/redundant-link-discovery.md; docs/three-node-soak.md'),
]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
def box(slide, x, y, w, h, text, size, color, bold=False):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, RGBColor.from_string(color)
    return tf

md = ['# Enlaces redundantes: módulos modificados y motivos', '',
      'Comparación reproducible: `git diff 8df7996 81aa146 -- components app tests`.',
      'Fuentes consultadas en el worktree `three-node-soak-20261015`. No incluye los cambios locales sin confirmar de `comnetar-stability-tools`.',
      'Los resultados de pruebas citados son los registrados en las fuentes; no se ejecutó firmware ni se repitieron esas pruebas al preparar esta presentación.', '']
for n, (title, subtitle, bullets, refs) in enumerate(slides, 1):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string('101D30')
    box(s, .6, .3, 12, .35, 'I4A  /  ARQUITECTURA Y CAMBIOS', 11, '51D5C5', True)
    box(s, .6, .9, 12, .65, title, 31, 'FFFFFF', True)
    box(s, .6, 1.65, 12, .6, subtitle, 19, '8DDCD4')
    for j, b in enumerate(bullets):
        box(s, .65, 2.5+j*.91, .4, .4, f'{j+1:02}', 13, '51D5C5', True)
        box(s, 1.2, 2.42+j*.91, 11.35, .88, b, 19, 'E4EBF5')
    box(s, .65, 6.65, 11.8, .56, 'Fuentes: '+refs, 9, '9BAEC7')
    box(s, 12.3, 6.85, .5, .3, str(n), 11, '51D5C5')
    s.notes_slide.notes_text_frame.text = 'Referencias: '+refs+'\nComparación: 8df7996 → 81aa146.\n'+'\n'.join(bullets)
    md += [f'## {n}. {title}', '', subtitle, ''] + ['- '+b for b in bullets] + ['', '**Referencias:** '+refs, '', '---', '']
prs.save(OUT / 'presentacion-enlaces-redundantes.pptx')
(OUT / 'presentacion-enlaces-redundantes.md').write_text('\n'.join(md), encoding='utf-8')
print(f'Generadas {len(slides)} diapositivas en {OUT}')
